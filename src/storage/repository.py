"""
Data-access layer — synchronous SQLAlchemy (using a thread-safe session factory).

All scraper operations are I/O-bound on the browser anyway, so sync DB calls
are perfectly fine and keep the code simple.
"""
from __future__ import annotations

import secrets
from contextlib import contextmanager
from datetime import datetime, timedelta
from typing import Generator

import bcrypt as _bcrypt_lib
from loguru import logger
from sqlalchemy import create_engine, select, text, update
from sqlalchemy.orm import Session, sessionmaker

from src.config import settings
from src.storage.models import (
    AppSettings, Base, JobLog, Lead, RunLog, ScheduledScan, ScrapeJob, User, UserSession,
)


def _hash_pw(password: str) -> str:
    return _bcrypt_lib.hashpw(password.encode(), _bcrypt_lib.gensalt()).decode()


def _verify_pw(password: str, hashed: str) -> bool:
    return _bcrypt_lib.checkpw(password.encode(), hashed.encode())


# ── Engine / session factory ──────────────────────────────────────────────────

_engine = create_engine(
    f"sqlite:///{settings.db_path}",
    connect_args={"check_same_thread": False},
    echo=False,
)

_SessionFactory = sessionmaker(bind=_engine, expire_on_commit=False)


def init_db() -> None:
    """Create all tables if they don't exist yet, then apply any pending migrations."""
    Base.metadata.create_all(_engine)
    _run_migrations()
    logger.info("Database ready: {}", settings.db_path)


def _run_migrations() -> None:
    """Add new columns to existing tables without data loss (SQLite ALTER TABLE)."""
    migrations = [
        "ALTER TABLE leads ADD COLUMN cronograma_json TEXT DEFAULT '[]'",
        "ALTER TABLE leads ADD COLUMN user_id INTEGER NOT NULL DEFAULT 0",
        "ALTER TABLE leads ADD COLUMN job_id INTEGER DEFAULT 0",
        "ALTER TABLE run_log ADD COLUMN user_id INTEGER NOT NULL DEFAULT 0",
        # scheduled_scans table — created via create_all; this is a no-op safety net
        (
            "CREATE TABLE IF NOT EXISTS scheduled_scans ("
            "id INTEGER PRIMARY KEY AUTOINCREMENT, "
            "name TEXT NOT NULL DEFAULT '', "
            "keywords_json TEXT NOT NULL DEFAULT '[]', "
            "use_ai BOOLEAN NOT NULL DEFAULT 0, "
            "download_pdf BOOLEAN NOT NULL DEFAULT 1, "
            "date_from TEXT NOT NULL DEFAULT '', "
            "date_to TEXT NOT NULL DEFAULT '', "
            "max_pages INTEGER NOT NULL DEFAULT 0, "
            "frequency TEXT NOT NULL DEFAULT 'daily', "
            "time_of_day TEXT NOT NULL DEFAULT '08:00', "
            "day_of_week INTEGER NOT NULL DEFAULT 0, "
            "enabled BOOLEAN NOT NULL DEFAULT 1, "
            "notify_keywords_json TEXT NOT NULL DEFAULT '[]', "
            "notify_emails_json TEXT NOT NULL DEFAULT '[]', "
            "created_at DATETIME DEFAULT CURRENT_TIMESTAMP, "
            "last_run_at DATETIME, "
            "next_run_at DATETIME)"
        ),
    ]
    with _engine.connect() as conn:
        for sql in migrations:
            try:
                conn.execute(text(sql))
                conn.commit()
            except Exception:
                # Column already exists — ignore
                pass


@contextmanager
def get_session() -> Generator[Session, None, None]:
    session = _SessionFactory()
    try:
        yield session
        session.commit()
    except Exception:
        session.rollback()
        raise
    finally:
        session.close()


# ── Users & auth ─────────────────────────────────────────────────────────────

def create_user(username: str, password: str, is_admin: bool = False) -> User:
    with get_session() as s:
        u = User(
            username=username,
            password_hash=_hash_pw(password),
            is_admin=is_admin,
        )
        s.add(u)
        s.flush()
        return u


def get_user_by_id(user_id: int) -> User | None:
    with get_session() as s:
        return s.get(User, user_id)


def get_user_by_username(username: str) -> User | None:
    with get_session() as s:
        return s.scalar(select(User).where(User.username == username))


def authenticate_user(username: str, password: str) -> User | None:
    user = get_user_by_username(username)
    if user and _verify_pw(password, user.password_hash):
        return user
    return None


def create_session(user_id: int, hours: int = 168) -> str:  # 7 days
    token = secrets.token_urlsafe(32)
    expires = datetime.utcnow() + timedelta(hours=hours)
    with get_session() as s:
        s.add(UserSession(token=token, user_id=user_id, expires_at=expires))
    return token


def get_session_user(token: str) -> User | None:
    """Return the User for this session token if it exists and hasn't expired."""
    with get_session() as s:
        sess = s.get(UserSession, token)
        if not sess or sess.expires_at < datetime.utcnow():
            return None
        return s.get(User, sess.user_id)


def delete_session(token: str) -> None:
    from sqlalchemy import delete as sql_delete
    with get_session() as s:
        s.execute(sql_delete(UserSession).where(UserSession.token == token))


def get_all_users() -> list[User]:
    with get_session() as s:
        return list(s.scalars(select(User).order_by(User.id)).all())


def delete_user(user_id: int) -> None:
    from sqlalchemy import delete as sql_delete
    with get_session() as s:
        s.execute(sql_delete(UserSession).where(UserSession.user_id == user_id))
        s.execute(sql_delete(User).where(User.id == user_id))


def ensure_admin_user() -> None:
    """Create a default admin account on first run if no users exist."""
    from sqlalchemy import func as sqlfunc
    with get_session() as s:
        count = s.scalar(select(sqlfunc.count()).select_from(User)) or 0
    if count == 0:
        create_user("admin", "admin1234", is_admin=True)
        logger.info("Default admin user created (username=admin, password=admin1234)")


# ── Run log ───────────────────────────────────────────────────────────────────

def create_run(run_id: str, user_id: int = 0) -> RunLog:
    with get_session() as s:
        run = RunLog(run_id=run_id, user_id=user_id, started_at=datetime.utcnow(), status="running")
        s.add(run)
    return run


def finish_run(
    run_id: str,
    rows_scanned: int,
    keyword_matches: int,
    leads_saved: int,
    status: str = "done",
    error_message: str = "",
) -> None:
    with get_session() as s:
        s.execute(
            update(RunLog)
            .where(RunLog.run_id == run_id)
            .values(
                finished_at=datetime.utcnow(),
                rows_scanned=rows_scanned,
                keyword_matches=keyword_matches,
                leads_saved=leads_saved,
                status=status,
                error_message=error_message,
            )
        )


# ── Leads ─────────────────────────────────────────────────────────────────────

def save_lead(lead: Lead) -> Lead:
    with get_session() as s:
        s.add(lead)
        s.flush()
        logger.info("Lead saved → id={} score={}", lead.id, lead.match_score)
    return lead


def get_leads(
    min_score: int = 0,
    reviewed: bool | None = None,
    limit: int = 200,
    offset: int = 0,
    user_id: int | None = None,
) -> list[Lead]:
    with get_session() as s:
        q = select(Lead).where(Lead.match_score >= min_score)
        if reviewed is not None:
            q = q.where(Lead.reviewed == reviewed)
        if user_id is not None:
            q = q.where(Lead.user_id == user_id)
        q = q.order_by(Lead.match_score.desc(), Lead.scraped_at.desc())
        q = q.limit(limit).offset(offset)
        return list(s.scalars(q).all())


def get_lead_by_id(lead_id: int) -> Lead | None:
    with get_session() as s:
        return s.get(Lead, lead_id)


def update_review(
    lead_id: int,
    reviewed: bool,
    review_notes: str = "",
    bid_decision: str = "",
) -> None:
    with get_session() as s:
        s.execute(
            update(Lead)
            .where(Lead.id == lead_id)
            .values(
                reviewed=reviewed,
                review_notes=review_notes,
                bid_decision=bid_decision,
            )
        )


def delete_all_leads(user_id: int | None = None) -> int:
    """Delete Lead rows. If user_id given, only that user's rows."""
    from sqlalchemy import delete as sql_delete
    with get_session() as s:
        q = sql_delete(Lead)
        if user_id is not None:
            q = q.where(Lead.user_id == user_id)
        result = s.execute(q)
        return result.rowcount


def count_leads(min_score: int = 0, user_id: int | None = None) -> int:
    from sqlalchemy import func as sqlfunc

    with get_session() as s:
        q = select(sqlfunc.count()).select_from(Lead).where(Lead.match_score >= min_score)
        if user_id is not None:
            q = q.where(Lead.user_id == user_id)
        return s.scalar(q) or 0


def get_recent_runs(limit: int = 10, user_id: int | None = None) -> list[RunLog]:
    with get_session() as s:
        q = select(RunLog).order_by(RunLog.started_at.desc())
        if user_id is not None:
            q = q.where(RunLog.user_id == user_id)
        q = q.limit(limit)
        return list(s.scalars(q).all())


# ── App Settings ─────────────────────────────────────────────────────────────

_DEFAULTS: dict[str, object] = {
    "date_from": "",   # "DD/MM/YYYY" – start of publication date range
    "date_to":   "",   # "DD/MM/YYYY" – end of publication date range
    "keywords": [
        "DATALOGGER",
        "MACROMEDIDOR",
        "MEDIDOR DE FLUJO",
        "CAUDALIMETRO ULTRASONICO",
        "MEDIDOR DE NIVEL",
        "TRANSMISOR DE FLUJO",
        "GEOFONO",
        "CORRELADOR",
        "SENSOR DE PRESION",
        "TRANSMISOR DE PRESION",
        "CALIBRADOR DE PRESION",
        "MULTICALIBRADOR DE PRESION",
        "CAMARA DE INSPECCION",
        "INDICADOR DE PRESION",
        "LOCALIZADOR DE AVERIAS",
        "LOCALIZADOR DE CABLES",
        "MANOMETRO",
        "TRANSDUCTOR DE PRESION",
        "MANIFOLD",
        "DETECTOR DE FUGA",
        "LOCALIZADOR DE FUGA",
        "DETECTOR DE METALES",
        "GEORADAR",
        "MULTIPARAMETRO",
        "MEDIDOR DE CLORO",
        "DETECTOR DE GAS",
        "EXPLOSIMETRO",
        "DETECTOR MULTIGAS",
        "CABINA DE FLUJO",
        "CABINA DE BIOSEGURIDAD",
        "MANTENIMIENTO DE MACROMEDIDORES",
        "CALIBRACION DE MEDIDORES",
        "CONTROL DE SECTORES IMPLEMENTADOS",
        "CALIBRACION",
        "INSTALACION DE MACROMEDIDORES",
        "DETECCION DE FUGAS",
        "MANTENIMIENTO",
    ],
    "spec_start_page": 20,
    "spec_end_page":   30,
    "max_pages":       0,   # 0 = unlimited (scan all pages)
}


def get_setting(key: str):
    """Return the parsed Python value for *key*, or the built-in default."""
    import json as _json
    with get_session() as s:
        row = s.get(AppSettings, key)
        if row is None:
            return _DEFAULTS.get(key)
        try:
            return _json.loads(row.value)
        except Exception:
            return row.value


def set_setting(key: str, value) -> None:
    """Persist *value* (any JSON-serialisable type) under *key*."""
    import json as _json
    from datetime import datetime as _dt
    serialised = _json.dumps(value, ensure_ascii=False)
    with get_session() as s:
        existing = s.get(AppSettings, key)
        if existing:
            existing.value = serialised
            existing.updated_at = _dt.utcnow()
        else:
            s.add(AppSettings(key=key, value=serialised))


def _user_setting_key(user_id: int, key: str) -> str:
    return f"u{user_id}.{key}"


def get_user_setting(user_id: int, key: str):
    """Return user-specific setting, falling back to global then default."""
    val = get_setting(_user_setting_key(user_id, key))
    if val is not None:
        return val
    return get_setting(key)


def set_user_setting(user_id: int, key: str, value) -> None:
    set_setting(_user_setting_key(user_id, key), value)


def get_all_user_settings(user_id: int) -> dict:
    """All settings for *user_id*, overriding globals with user-specific values."""
    import json as _json
    result = get_all_settings()  # start with global + defaults
    # Strip per-user keys from the result (they appear as "u42.keywords" etc.)
    prefix = f"u{user_id}."
    with get_session() as s:
        for row in s.scalars(
            select(AppSettings).where(AppSettings.key.like(f"{prefix}%"))
        ).all():
            short_key = row.key[len(prefix):]
            try:
                result[short_key] = _json.loads(row.value)
            except Exception:
                result[short_key] = row.value
    # Remove keys from other users that leaked in via get_all_settings
    result = {k: v for k, v in result.items() if not k.startswith("u")}
    return result


def get_all_settings() -> dict:
    """Return all settings merged with defaults (DB values override defaults)."""
    import json as _json
    result = dict(_DEFAULTS)
    with get_session() as s:
        for row in s.scalars(select(AppSettings)).all():
            try:
                result[row.key] = _json.loads(row.value)
            except Exception:
                result[row.key] = row.value
    return result


# ── Excel export ──────────────────────────────────────────────────────────────

def export_to_excel(path: str, min_score: int = 0, user_id: int | None = None) -> str:
    """Export leads to Excel. If user_id given, only that user's rows."""
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill

    leads = get_leads(min_score=min_score, limit=10_000, user_id=user_id)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "SEACE Leads"

    headers = [
        "ID", "Fecha Scraping", "Entidad", "Nomenclatura", "Tipo Objeto",
        "Descripción", "Score", "Nivel", "Resumen IA", "URL Ficha",
        "PDF", "Revisado", "Decisión", "Notas",
    ]

    # Header row style
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True)
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    # Score colour bands
    def _row_fill(score: int) -> PatternFill:
        if score >= 75:
            return PatternFill("solid", fgColor="C6EFCE")   # green
        if score >= 40:
            return PatternFill("solid", fgColor="FFEB9C")   # yellow
        return PatternFill("solid", fgColor="FFC7CE")        # red

    for row_idx, lead in enumerate(leads, 2):
        values = [
            lead.id,
            lead.scraped_at.strftime("%Y-%m-%d %H:%M") if lead.scraped_at else "",
            lead.entity,
            lead.nomenclature,
            lead.object_type,
            lead.description,
            lead.match_score,
            lead.match_level,
            lead.ai_summary,
            lead.ficha_url,
            lead.pdf_local_path,
            "Si" if lead.reviewed else "No",
            lead.bid_decision,
            lead.review_notes,
        ]
        fill = _row_fill(lead.match_score)
        for col_idx, value in enumerate(values, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.fill = fill
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    # Column widths
    widths = [6, 16, 35, 30, 20, 50, 7, 10, 60, 40, 40, 9, 10, 40]
    for col_idx, width in enumerate(widths, 1):
        ws.column_dimensions[
            openpyxl.utils.get_column_letter(col_idx)
        ].width = width

    wb.save(path)
    logger.info("Excel exported: {}", path)
    return path


# ── ScrapeJob ──────────────────────────────────────────────────────────────────

def create_scrape_job(
    user_id: int,
    keywords: list[str],
    date_from: str = "",
    date_to: str = "",
    search_query: str = "",
    download_pdf: bool = True,
    use_ai: bool = True,
    max_pages: int = 0,
) -> ScrapeJob:
    import json as _json
    job = ScrapeJob(
        user_id=user_id,
        keywords_json=_json.dumps(keywords, ensure_ascii=False),
        date_from=date_from,
        date_to=date_to,
        search_query=search_query,
        download_pdf=download_pdf,
        use_ai=use_ai,
        max_pages=max_pages,
        status="pending",
    )
    with get_session() as s:
        s.add(job)
        s.flush()
        job_id = job.id
    with get_session() as s:
        return s.get(ScrapeJob, job_id)


def get_scrape_job(job_id: int) -> ScrapeJob | None:
    with get_session() as s:
        return s.get(ScrapeJob, job_id)


def list_scrape_jobs(user_id: int | None = None, limit: int = 50) -> list[ScrapeJob]:
    with get_session() as s:
        q = select(ScrapeJob).order_by(ScrapeJob.created_at.desc()).limit(limit)
        if user_id is not None:
            q = q.where(ScrapeJob.user_id == user_id)
        return list(s.scalars(q).all())


def update_scrape_job(job_id: int, **kwargs) -> None:
    with get_session() as s:
        s.execute(update(ScrapeJob).where(ScrapeJob.id == job_id).values(**kwargs))


# ── JobLog ─────────────────────────────────────────────────────────────────────

def add_job_log(job_id: int, level: str, message: str) -> None:
    from datetime import datetime as _dt
    with get_session() as s:
        s.add(JobLog(job_id=job_id, level=level.upper(), message=message,
                     ts=_dt.utcnow()))


def get_job_logs(job_id: int, since_id: int = 0, limit: int = 500) -> list[JobLog]:
    with get_session() as s:
        q = (
            select(JobLog)
            .where(JobLog.job_id == job_id, JobLog.id > since_id)
            .order_by(JobLog.id)
            .limit(limit)
        )
        return list(s.scalars(q).all())


# ── Leads by run ───────────────────────────────────────────────────────────────

def get_leads_by_run_id(run_id: str) -> list[Lead]:
    """Return all leads saved during a specific scraper run."""
    with get_session() as s:
        return list(s.scalars(select(Lead).where(Lead.run_id == run_id)).all())


# ── ScheduledScan ──────────────────────────────────────────────────────────────

def create_scheduled_scan(
    name: str,
    keywords: list[str],
    frequency: str = "daily",
    time_of_day: str = "08:00",
    day_of_week: int = 0,
    enabled: bool = True,
    use_ai: bool = False,
    download_pdf: bool = True,
    date_from: str = "",
    date_to: str = "",
    max_pages: int = 0,
    notify_keywords: list[str] | None = None,
    notify_emails: list[str] | None = None,
) -> ScheduledScan:
    import json as _json
    scan = ScheduledScan(
        name=name,
        keywords_json=_json.dumps(keywords, ensure_ascii=False),
        frequency=frequency,
        time_of_day=time_of_day,
        day_of_week=day_of_week,
        enabled=enabled,
        use_ai=use_ai,
        download_pdf=download_pdf,
        date_from=date_from,
        date_to=date_to,
        max_pages=max_pages,
        notify_keywords_json=_json.dumps(notify_keywords or [], ensure_ascii=False),
        notify_emails_json=_json.dumps(notify_emails or [], ensure_ascii=False),
    )
    with get_session() as s:
        s.add(scan)
        s.flush()
        scan_id = scan.id
    with get_session() as s:
        return s.get(ScheduledScan, scan_id)


def get_scheduled_scan(scan_id: int) -> ScheduledScan | None:
    with get_session() as s:
        return s.get(ScheduledScan, scan_id)


def list_scheduled_scans(enabled_only: bool = False) -> list[ScheduledScan]:
    with get_session() as s:
        q = select(ScheduledScan).order_by(ScheduledScan.id)
        if enabled_only:
            q = q.where(ScheduledScan.enabled == True)  # noqa: E712
        return list(s.scalars(q).all())


def update_scheduled_scan(scan_id: int, **kwargs) -> None:
    with get_session() as s:
        s.execute(
            update(ScheduledScan).where(ScheduledScan.id == scan_id).values(**kwargs)
        )


def delete_scheduled_scan(scan_id: int) -> None:
    from sqlalchemy import delete as sql_delete
    with get_session() as s:
        s.execute(sql_delete(ScheduledScan).where(ScheduledScan.id == scan_id))
